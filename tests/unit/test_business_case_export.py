"""The business case Word and PDF exports.

Each writer's bytes are opened with pypdf / python-docx and the title, a canvas
entry, the status and a figure are found in the text. A second organisation's
case answers 404 on every format and none of its text appears in a document.
"""

import io
import uuid
from decimal import Decimal

import pytest

FORMATS = ("docx", "pdf", "mermaid", "lucid", "archi")


def _make_user(db_session, org, label):
    from app.models.user import Role, User
    from tests.smoke.conftest import PASSWORD

    Role.insert_roles()
    user = User(
        email=f"export-{label}-{uuid.uuid4().hex[:8]}@example.com",
        first_name="Export", last_name=label, organization_id=org.id,
        enterprise_role="business_architect", confirmed=True,
    )
    user.role = Role.query.filter_by(name="Architect").one()
    user.password = PASSWORD
    db_session.add(user)
    db_session.flush()
    return user


def _make_case(db_session, org, user, title, **fields):
    from app.models.business_case import BusinessCase

    case = BusinessCase(title=title, organization_id=org.id, created_by_id=user.id, **fields)
    db_session.add(case)
    db_session.flush()
    return case


@pytest.fixture
def two_orgs(db_session, make_org):
    suffix = uuid.uuid4().hex[:6]
    mine, other = make_org("mine"), make_org("other")
    me, them = _make_user(db_session, mine, "mine"), _make_user(db_session, other, "other")
    case = _make_case(
        db_session, mine, me, f"Consolidate regional CRM {suffix}",
        description="One platform replaces the diverged regional instances.",
        status="submitted",
        problem_statement="Three regional CRM instances disagree on customer counts.",
        recommended_option="Adopt the unified platform in two waves.",
        capex=Decimal("123456.78"), roi_percentage=Decimal("17.50"),
    )
    rival = _make_case(
        db_session, other, them, f"Rival secret programme {suffix}",
        description="Rival confidential description.",
        problem_statement="Rival confidential problem.",
        capex=Decimal("987654.32"),
    )
    return me, case, rival


def _context(case):
    from app.modules.business_case import service

    return service.build_export_context(case)


def _pdf_text(body):
    from pypdf import PdfReader

    return " ".join(page.extract_text() for page in PdfReader(io.BytesIO(body)).pages)


def _docx_text(body):
    from docx import Document

    document = Document(io.BytesIO(body))
    cells = [c.text for t in document.tables for r in t.rows for c in r.cells]
    return "\n".join([p.text for p in document.paragraphs] + cells)


def test_pdf_carries_title_status_entry_and_figure(app, two_orgs):
    from app.services.togaf_deliverable_export_service import TOGAFDeliverableExportService

    _me, case, _rival = two_orgs
    body = TOGAFDeliverableExportService().export_business_case_to_pdf(case, _context(case))
    assert body.startswith(b"%PDF")
    text = _pdf_text(body)
    assert case.title in text
    assert "Submitted" in text
    assert "Three regional CRM instances disagree" in text
    assert "123456.78" in text


def test_docx_carries_title_status_entry_and_figure(app, two_orgs):
    from app.services.togaf_deliverable_export_service import TOGAFDeliverableExportService

    _me, case, _rival = two_orgs
    body = TOGAFDeliverableExportService().export_business_case_to_docx(case, _context(case))
    assert body.startswith(b"PK")
    text = _docx_text(body)
    assert case.title in text
    assert "Submitted" in text
    assert "Three regional CRM instances disagree" in text
    assert "123456.78" in text
    assert "17.50" in text


def test_empty_boxes_and_unset_figures_are_stated_not_invented(app, db_session, make_org):
    from app.services.togaf_deliverable_export_service import TOGAFDeliverableExportService

    org = make_org("bare")
    case = _make_case(db_session, org, _make_user(db_session, org, "bare"), "Bare case")
    text = _docx_text(TOGAFDeliverableExportService().export_business_case_to_docx(case, _context(case)))
    assert "Nothing recorded yet." in text
    assert "Not linked" in text
    assert "—" in text


@pytest.mark.parametrize("fmt", FORMATS)
def test_other_organisations_case_answers_404_on_every_format(
    app, client, login_as, two_orgs, fmt
):
    me, _case, rival = two_orgs
    login_as(client, me)
    response = client.get(f"/business-case/{rival.id}/export?format={fmt}")
    assert response.status_code == 404
    assert b"Rival" not in response.data


@pytest.mark.parametrize("fmt,magic", [("docx", b"PK"), ("pdf", b"%PDF")])
def test_route_serves_own_case_and_nothing_of_the_other(
    app, client, login_as, two_orgs, fmt, magic
):
    me, case, _rival = two_orgs
    login_as(client, me)
    response = client.get(f"/business-case/{case.id}/export?format={fmt}")
    assert response.status_code == 200
    assert response.data.startswith(magic)
    assert fmt in response.headers["Content-Disposition"]
    assert response.headers["Content-Disposition"].startswith("attachment; filename=")
    text = _pdf_text(response.data) if fmt == "pdf" else _docx_text(response.data)
    assert case.title in text
    for leaked in ("Rival", "987654.32"):
        assert leaked not in text
